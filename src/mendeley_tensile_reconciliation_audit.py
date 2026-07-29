from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = (
    PROJECT_ROOT / "data/public/mendeley-8jxw5533py-v2"
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT / "results/mendeley_tensile_reconciliation_audit"
)
RAW_WORKBOOK = SOURCE_ROOT / "source/3-2_Original_data_of_material_tests.xlsx"
SUMMARY_WORKBOOK = SOURCE_ROOT / "source/3-1_Material_property.xlsx"
SOURCE_MANIFEST = SOURCE_ROOT / "manifest.csv"

EXPECTED_HASHES = {
    RAW_WORKBOOK: (
        "0e9ef3c7ae638cb3544c119f21531a7f963f91ffc490976bd21ae41bb80724cc"
    ),
    SUMMARY_WORKBOOK: (
        "bc59ed28723893f111c2cf5e3cd8f70c8f51ea627f28a232501cd13557c8fbcd"
    ),
}
SUMMARY_HEADER_ROWS = (1, 4, 7, 10, 13, 16, 19)
PEAK_MATCH_ABS_TOLERANCE_MPA = 5.0e-5


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _normalize_label(label: str) -> str:
    return label.replace("45-", "45Deg-")


def read_literal_raw_peaks(
    manifest_path: Path = SOURCE_MANIFEST,
) -> List[Dict[str, object]]:
    with Path(manifest_path).open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    result = []
    for row in rows:
        result.append(
            {
                "raw_label": row["source_specimen_label"],
                "raw_peak_stress_MPa": float(
                    row["extracted_peak_stress_MPa"]
                ),
                "source_platform_axis": row["source_platform_axis"],
                "source_placement": row["source_placement"],
                "curve_file": row["file"],
            }
        )
    return result


def read_summary_specimen_cells(
    workbook_path: Path = SUMMARY_WORKBOOK,
) -> List[Dict[str, object]]:
    workbook = load_workbook(
        filename=Path(workbook_path), read_only=True, data_only=True
    )
    sheet = workbook[workbook.sheetnames[0]]
    rows = []
    for header_row in SUMMARY_HEADER_ROWS:
        group = str(sheet.cell(header_row, 1).value)
        for column in (2, 3, 4):
            label = _normalize_label(
                str(sheet.cell(header_row, column).value)
            )
            rows.append(
                {
                    "summary_group": group,
                    "summary_label": label,
                    "summary_UTS_MPa": float(
                        sheet.cell(header_row + 1, column).value
                    ),
                    "summary_modulus_MPa": float(
                        str(
                            sheet.cell(
                                header_row + 2, column
                            ).value
                        ).rstrip("`")
                    ),
                }
            )
    workbook.close()
    return rows


def reconcile_summary_to_raw_peaks(
    raw_rows: List[Dict[str, object]],
    summary_rows: List[Dict[str, object]],
    tolerance_MPa: float = PEAK_MATCH_ABS_TOLERANCE_MPA,
) -> List[Dict[str, object]]:
    mappings = []
    for summary in summary_rows:
        value = float(summary["summary_UTS_MPa"])
        candidates = [
            raw
            for raw in raw_rows
            if abs(float(raw["raw_peak_stress_MPa"]) - value)
            <= tolerance_MPa
        ]
        labels = sorted(str(row["raw_label"]) for row in candidates)
        same_label = str(summary["summary_label"]) in labels
        mappings.append(
            {
                **summary,
                "matching_raw_curve_count": len(candidates),
                "matching_raw_labels": " | ".join(labels),
                "same_label_peak_match": same_label,
                "cross_label_or_missing_match": not same_label,
            }
        )
    return mappings


def write_reconciliation_audit(
    output_dir: Path = DEFAULT_OUTPUT,
) -> Dict[str, object]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    actual_hashes = {
        str(path.relative_to(PROJECT_ROOT)): _sha256(path)
        for path in EXPECTED_HASHES
    }
    checksums_passed = all(
        actual_hashes[str(path.relative_to(PROJECT_ROOT))] == expected
        for path, expected in EXPECTED_HASHES.items()
    )
    if not checksums_passed:
        raise ValueError("Mendeley tensile source checksum mismatch")

    raw_rows = read_literal_raw_peaks()
    summary_rows = read_summary_specimen_cells()
    mappings = reconcile_summary_to_raw_peaks(raw_rows, summary_rows)

    matched_labels = [
        label
        for row in mappings
        for label in str(row["matching_raw_labels"]).split(" | ")
        if label
    ]
    match_use_counts = Counter(matched_labels)
    raw_labels = {str(row["raw_label"]) for row in raw_rows}
    used_raw_labels = set(match_use_counts)
    omitted_raw_labels = sorted(raw_labels - used_raw_labels)
    repeated_raw_labels = {
        label: count
        for label, count in sorted(match_use_counts.items())
        if count > 1
    }
    literal_group_counts = Counter(
        (
            str(row["source_platform_axis"]),
            str(row["source_placement"]),
        )
        for row in raw_rows
    )

    same_label_matches = sum(
        bool(row["same_label_peak_match"]) for row in mappings
    )
    all_summary_peaks_found = all(
        int(row["matching_raw_curve_count"]) >= 1 for row in mappings
    )
    one_to_one = (
        len(raw_rows) == len(summary_rows)
        and len(used_raw_labels) == len(raw_rows)
        and not repeated_raw_labels
        and all(
            int(row["matching_raw_curve_count"]) == 1
            and bool(row["same_label_peak_match"])
            for row in mappings
        )
    )

    mapping_path = output_dir / "summary_to_raw_peak_mapping.csv"
    with mapping_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(mappings[0]))
        writer.writeheader()
        writer.writerows(mappings)

    group_rows = [
        {
            "source_platform_axis": axis,
            "source_placement": placement,
            "literal_curve_count": count,
            "expected_replicates": 3,
            "count_matches_article_claim": count == 3,
        }
        for (axis, placement), count in sorted(literal_group_counts.items())
    ]
    group_path = output_dir / "literal_group_counts.csv"
    with group_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(group_rows[0]))
        writer.writeheader()
        writer.writerows(group_rows)

    payload = {
        "audit_id": "mendeley_8jxw5533py_tensile_reconciliation_v1",
        "status": "source_reconciliation_failed_non_bijective",
        "material": "Bambu PLA Basic",
        "source_dataset_doi": "10.17632/8jxw5533py.2",
        "source_article_doi": "10.1016/j.dib.2026.113015",
        "source_license": "CC BY 4.0",
        "source_hashes": actual_hashes,
        "counts": {
            "article_claimed_specimens": 21,
            "literal_raw_curve_blocks": len(raw_rows),
            "summary_specimen_cells": len(summary_rows),
            "summary_cells_with_a_raw_peak_match": sum(
                int(row["matching_raw_curve_count"]) >= 1
                for row in mappings
            ),
            "same_label_peak_matches": same_label_matches,
            "unique_raw_curves_referenced_by_summary": len(
                used_raw_labels
            ),
            "raw_curves_omitted_from_summary_peak_values": len(
                omitted_raw_labels
            ),
            "raw_curves_reused_by_multiple_summary_cells": len(
                repeated_raw_labels
            ),
        },
        "omitted_raw_labels": omitted_raw_labels,
        "repeated_raw_labels": repeated_raw_labels,
        "literal_group_counts": {
            f"{axis}-{placement}": count
            for (axis, placement), count in sorted(
                literal_group_counts.items()
            )
        },
        "gates": {
            "source_checksums_passed": checksums_passed,
            "all_summary_peaks_found_in_literal_raw_curves": (
                all_summary_peaks_found
            ),
            "summary_to_raw_mapping_is_one_to_one": one_to_one,
            "all_summary_labels_match_same_label_raw_peaks": (
                same_label_matches == len(summary_rows)
            ),
            "all_literal_groups_have_three_replicates": all(
                count == 3 for count in literal_group_counts.values()
            ),
            "directional_calibration_release_allowed": False,
        },
        "decision": (
            "The summary workbook is not a one-to-one inventory of the raw "
            "curves: 21 summary cells reuse only 13 of 22 literal curves, "
            "nine raw curves are omitted, and several summary cells point "
            "to a differently labelled raw peak. The source cannot be "
            "relabelled or partitioned for directional calibration without "
            "author clarification."
        ),
        "allowed_use": (
            "Pooled, robust isotropic small-strain screening and negative "
            "data-quality evidence only. Do not use these labels to fit or "
            "validate local material axes."
        ),
        "outputs": {
            "summary_to_raw_peak_mapping": str(mapping_path),
            "literal_group_counts": str(group_path),
        },
    }
    provenance_path = output_dir / "provenance.json"
    provenance_path.write_text(
        json.dumps(payload, indent=2) + "\n", encoding="utf-8"
    )
    report = f"""# Mendeley Bambu PLA tensile reconciliation audit

Status: **{payload['status']}**

The immutable raw and summary workbook checksums pass. The article claims 21
specimens, while the raw workbook contains {len(raw_rows)} literal curve
blocks and the summary contains {len(summary_rows)} specimen cells.

Peak-value reconciliation is non-bijective:

- {same_label_matches}/{len(summary_rows)} summary cells match the raw curve
  carrying the same label.
- The 21 summary cells reference only {len(used_raw_labels)} unique raw
  curves.
- {len(omitted_raw_labels)} literal raw curves never appear among the summary
  peak values.
- {len(repeated_raw_labels)} raw curves are reused by multiple summary cells.

This is a source-data failure, not missing-data imputation. No automatic
relabeling, directional partition, Poisson inference, or local-axis material
card is released. The literal curves remain useful only for robust pooled
isotropic screening.
"""
    (output_dir / "report.md").write_text(report, encoding="utf-8")
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Audit the exact-brand tensile workbook reconciliation"
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    payload = write_reconciliation_audit(args.output)
    counts = payload["counts"]
    print(
        f"{payload['status']}; same-label matches="
        f"{counts['same_label_peak_matches']}/"
        f"{counts['summary_specimen_cells']}; unique raw used="
        f"{counts['unique_raw_curves_referenced_by_summary']}/"
        f"{counts['literal_raw_curve_blocks']}"
    )


if __name__ == "__main__":
    main()
